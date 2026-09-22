import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch
from django.test import TestCase
from openpyxl import Workbook, load_workbook
import pandas as pd

from .processors.normalize_manager import process_manager_sheet, parse_metric_type
from .processors.snapshot_engine import merge_plans_with_1c
from .processors.export_manager_facts import build_fact_lookup, update_manager_workbook
from .processors.combined_report import export_combined_report, combined_is_current


def source_book(path, client='Client'):
    w = Workbook()
    s = w.active
    s.append(['', '', '', '', '', '', '2026-I', '2026-II'])
    s.append(['Клиент', 'Класс товара', 'Номер изделия/ артикул', 'Производственный индекс',
              'КОД из 1С', 'Рабочее наименование из 1 С', 'Цена, юань', 'Цена, юань',
              'Январь 2026', '', '', '', 'Сентябрь 2026', '', '', ''])
    s.append(['', '', '', '', '', '', '', '', 'AOP, шт', 'Прогноз, шт', 'Факт, шт', 'AOP негатив, шт',
              'AOP, шт', 'Прогноз, шт', 'Факт, шт', 'Комментарий'])
    s.append([client, 'Class', 'A1', 'Index', 'Internal code', 'Working name', 10, 20, 2, 3, 8, 999, 4, 5, 9, 'Keep'])
    s['Q4'] = '=SUM(K4,O4)'
    w.save(path)
    w.close()


def plan_rows():
    return pd.DataFrame([{'Клиент': 'Client', 'Артикул': 'A1', 'Менеджер': 'Редько Вадим',
        'Код товара': 'Internal code',
        'Год': 2026, 'Номер месяца': month, 'AOP, шт': 2, 'Прогноз, шт': 3,
        'Цена, юань, без НДС 1 п/г 2026': 10, 'Класс товара': 'Class', 'Производственный индекс': 'Index'}
        for month in [1, 9]])


class NewFormatTests(TestCase):
    def test_plans_only_snapshot_is_allowed(self):
        from .processors.snapshot_engine import create_full_snapshot
        plans = plan_rows().iloc[[1]].copy()
        with tempfile.TemporaryDirectory() as folder:
            raw = Path(folder) / 'raw'
            raw.mkdir()
            with patch('uploads.processors.snapshot_engine.normalize_all_managers', return_value=plans), \
                 patch('uploads.processors.snapshot_engine.export_all_manager_fact_files', return_value={}), \
                 patch('uploads.processors.snapshot_engine.export_combined_report', return_value={'rows': 0}):
                result = create_full_snapshot(str(raw), 'plans_only')
        self.assertEqual(len(result), 1)
        self.assertEqual(result['Факт, шт'].sum(), 0)
        self.assertEqual(result.attrs['processing_info']['report_period'], '2026-09')

    def test_diagnostics_show_article_candidates_without_matching_wrong_code(self):
        plans = plan_rows().iloc[[1]].copy()
        plans['_Исходная строка'] = 18
        actual = pd.DataFrame([{'Клиент': 'Client', 'Артикул': 'A1', 'Код товара': 'Wrong code',
                               'Год': 2026, 'Номер месяца': 9, 'Факт, шт': 7}])
        result = merge_plans_with_1c(plans, actual)
        item = result.attrs['matching_report']['diagnostics'][0]
        self.assertEqual(result['Факт, шт'].sum(), 0)
        self.assertEqual(item['product_code'], 'Wrong code')
        self.assertEqual(item['candidates'][0]['product_code'], 'Internal code')
        self.assertEqual(item['candidates'][0]['price'], 10)
        self.assertEqual(item['candidates'][0]['forecast_quantity'], 3)
        self.assertEqual(item['candidates'][0]['source_row'], 18)

    def test_missing_price_is_in_diagnostics_without_changing_match_count(self):
        plans = plan_rows().iloc[[1]].copy()
        plans['Цена, юань, без НДС 1 п/г 2026'] = 0
        actual = pd.DataFrame([{'Клиент': 'Client', 'Артикул': 'A1',
                               'Год': 2026, 'Номер месяца': 9, 'Факт, шт': 7}])
        result = merge_plans_with_1c(plans, actual)
        report = result.attrs['matching_report']
        self.assertEqual(report['matched_rows'], 1)
        self.assertEqual(report['diagnostics'][0]['reason'], 'Не указана цена за период')
        self.assertIsNone(report['diagnostics'][0]['candidates'][0]['price'])

    def test_pool_clients_come_from_uploaded_list_not_substrings(self):
        plans = plan_rows().iloc[[1]].copy()
        plans['Менеджер'] = 'Ушаков Алексей'
        plans['_Исходные клиенты'] = 'ООО "ТД КАМАЦЕНТРСЕРВИС";\nТорговый дом Тонар; НОВЫЙ КЛИЕНТ ООО'
        actual = pd.DataFrame([{'Клиент': client, 'Артикул': 'A1', 'Год': 2026,
                               'Номер месяца': 9, 'Факт, шт': 2}
                              for client in ['ТД КАМАЦЕНТРСЕРВИС ООО', 'Торговый дом Тонар',
                                             'НОВЫЙ КЛИЕНТ ООО', 'ТОНАР ООО МЗ', 'Норма посторонняя ООО']])
        result = merge_plans_with_1c(plans, actual)
        self.assertEqual(result['Факт, шт'].sum(), 6)
        unmatched = result.attrs['matching_report']['unmatched_df']
        self.assertEqual(set(unmatched['Клиент из выгрузки']), {'ТОНАР ООО МЗ', 'Норма посторонняя ООО'})

    def test_unmatched_report_locates_managers_without_inventing_owner(self):
        plans = plan_rows().iloc[[1]].copy()
        plans['_Исходный файл'] = 'redko.xlsx'
        plans['_Исходный лист'] = 'План'
        plans['_Исходная строка'] = 12
        second = plans.copy()
        second['Менеджер'] = 'Другой менеджер'
        second['_Исходный файл'] = 'other.xlsx'
        plans = pd.concat([plans, second], ignore_index=True)
        actual = pd.DataFrame([{'Клиент': client, 'Артикул': article, 'Год': 2026,
                               'Номер месяца': 9, 'Факт, шт': 2}
                              for client, article in [('Client', 'A1'), ('Client', 'Absent'), ('Unknown', 'A1')]])
        report = merge_plans_with_1c(plans, actual).attrs['matching_report']['unmatched_df']
        known = report[report['Клиент из выгрузки'] == 'Client']
        self.assertTrue(known['Менеджер'].str.contains('Редько Вадим').all())
        self.assertTrue(known['Менеджер'].str.contains('Другой менеджер').all())
        self.assertTrue(known['Файл менеджера'].str.contains('redko.xlsx', regex=False).all())
        self.assertIn('строка 12', known[known['Артикул из выгрузки'] == 'A1'].iloc[0]['Где искать'])
        unknown = report[report['Клиент из выгрузки'] == 'Unknown'].iloc[0]
        self.assertEqual(unknown['Менеджер'], 'Не определён: клиент не найден')
        self.assertEqual(unknown['Файл менеджера'], '')

    def test_downloaded_workbook_is_next_runs_history_without_archive(self):
        from .processors.snapshot_engine import create_full_snapshot
        with tempfile.TemporaryDirectory() as folder:
            raw = Path(folder) / 'raw'
            raw.mkdir()
            source = raw / 'plan_redko_source.xlsx'
            source_book(source)
            (raw / 'fact_1c_test.xlsx').touch()
            archive = Path(folder) / 'processed' / 'final' / 'old'
            archive.mkdir(parents=True)
            (archive / 'FINAL_SALES_FACT_TABLE.xlsx').write_bytes(b'Archive must not be read')
            actual = pd.DataFrame([{'Клиент': 'Client', 'Артикул': 'A1',
                'Год': 2026, 'Номер месяца': 9, 'Факт, шт': 7}])
            actual.attrs['report_period'] = '2026-09'
            with patch('uploads.processors.snapshot_engine.normalize_1c_file', return_value=actual), patch(
                'uploads.processors.snapshot_engine.extract_existing_facts', side_effect=AssertionError('Archive read')
            ):
                first = create_full_snapshot(str(raw), 'first')
                self.assertEqual(first['Факт, шт'].tolist(), [8, 7])
                self.assertEqual(first['Факт, CNY'].tolist(), [80, 140])
                self.assertEqual(first.attrs['matching_report']['total_rows'], 1)
                downloaded = Path(first.attrs['processing_info']['manager_reports']['redko'])
                source.write_bytes(downloaded.read_bytes())
                actual['Номер месяца'] = 1
                actual['Факт, шт'] = 4
                actual.attrs['report_period'] = '2026-01'
                second = create_full_snapshot(str(raw), 'second')
                self.assertEqual(second['Факт, шт'].tolist(), [4, 7])
                self.assertEqual(second['Факт, CNY'].tolist(), [40, 140])
                repeated = create_full_snapshot(str(raw), 'third')
                self.assertEqual(repeated['Факт, шт'].tolist(), [4, 7])
                self.assertEqual(second.attrs['matching_report']['history_conflicts'], [])
                # Release Excel reader handles before Windows removes the fixture.
                import gc
                gc.collect()

    def test_current_month_without_match_is_reset_other_month_keeps_fact(self):
        plans = plan_rows()
        plans['Факт, шт'] = [8, 999]
        actual = pd.DataFrame([{'Клиент': 'Unknown', 'Артикул': 'A1',
                               'Год': 2026, 'Номер месяца': 9, 'Факт, шт': 7}])
        result = merge_plans_with_1c(plans, actual)
        self.assertEqual(result['Факт, шт'].tolist(), [8, 0])
        self.assertEqual(result['Факт, CNY'].tolist(), [80, 0])
        self.assertEqual(result.attrs['matching_report']['unmatched_rows'], 1)

    def test_revenue_is_ignored_and_quantity_uses_period_price(self):
        plans = plan_rows()
        plans['Факт, шт'] = [8, 999]
        plans['Факт, CNY'] = [123, 9990]
        plans.loc[plans['Номер месяца'] == 9, 'Цена, юань, без НДС 1 п/г 2026'] = 20
        actual = pd.DataFrame([{'Клиент': 'Client', 'Артикул': 'A1', 'Год': 2026, 'Номер месяца': 9,
                              'Факт, шт': 7, 'Факт, CNY': 99999999}])
        history = pd.DataFrame([{'_key_client': 'client', '_key_article': 'a1', '_key_month': '2026-01',
                                'Факт, шт': 8, 'Факт, CNY': 123}])
        result = merge_plans_with_1c(plans, actual, history)
        self.assertEqual(result['Факт, CNY'].tolist(), [123, 140])
        self.assertEqual(result['Факт, шт'].tolist(), [8, 7])

    def test_missing_price_keeps_quantity_and_blank_money(self):
        plans = plan_rows().iloc[[1]].copy()
        plans['Цена, юань, без НДС 1 п/г 2026'] = 0
        actual = pd.DataFrame([{'Клиент': 'Client', 'Артикул': 'A1', 'Год': 2026, 'Номер месяца': 9, 'Факт, шт': 7}])
        result = merge_plans_with_1c(plans, actual)
        self.assertEqual(result.iloc[0]['Факт, шт'], 7)
        self.assertTrue(pd.isna(result.iloc[0]['Факт, CNY']))
        self.assertEqual(result.attrs['matching_report']['missing_price_rows'], 1)

    def test_price_half_name_and_negative_scenario(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'plan_redko_test.xlsx'
            source_book(path)
            out = process_manager_sheet(pd.read_excel(path, header=None), filename=path.name)
        self.assertEqual(len(out), 2)
        self.assertEqual(out.iloc[0]['Наименование'], 'Working name')
        self.assertEqual(out['AOP, шт'].tolist(), [2, 4])
        self.assertEqual(out['Цена, юань, без НДС 1 п/г 2026'].tolist(), [10, 20])
        self.assertEqual(out.iloc[0]['Класс товара'], 'Class')
        self.assertEqual(out.iloc[0]['Производственный индекс'], 'Index')
        self.assertEqual(out.iloc[0]['Код товара'], 'Internal code')

    def test_product_code_is_primary_and_article_is_fallback(self):
        plans = plan_rows().iloc[[1]].copy()
        plans['Код товара'] = 'CODE-1'
        actual = pd.DataFrame([{
            'Клиент': 'Client', 'Артикул': 'Изменённый артикул', 'Код товара': 'CODE-1',
            'Год': 2026, 'Номер месяца': 9, 'Факт, шт': 7,
        }])
        result = merge_plans_with_1c(plans, actual)
        self.assertEqual(result.iloc[0]['Факт, шт'], 7)
        self.assertEqual(result.attrs['matching_report']['matched_rows'], 1)

        old_actual = actual.drop(columns=['Код товара']).copy()
        old_actual['Артикул'] = 'A1'
        result = merge_plans_with_1c(plans, old_actual)
        self.assertEqual(result.iloc[0]['Факт, шт'], 7)

    def test_exact_client_wins_before_short_name_fallback(self):
        plans = plan_rows().iloc[[1]].copy()
        plans['Код товара'] = 'CODE-1'
        second = plans.copy()
        plans['Клиент'] = 'Клиент ООО'
        second['Клиент'] = 'Клиент АО'
        plans = pd.concat([plans, second], ignore_index=True)
        actual = pd.DataFrame([{
            'Клиент': 'Клиент ООО', 'Артикул': 'A1', 'Код товара': 'CODE-1',
            'Год': 2026, 'Номер месяца': 9, 'Факт, шт': 7,
        }])
        result = merge_plans_with_1c(plans, actual)
        self.assertEqual(result['Факт, шт'].tolist(), [7, 0])

    def test_bilingual_client_uses_short_name_fallback(self):
        plans = plan_rows().iloc[[1]].copy()
        plans['Клиент'] = 'КЛИЕНТ А ООО / CLIENT A LLC'
        plans['Код товара'] = 'CODE-1'
        actual = pd.DataFrame([{
            'Клиент': 'КЛИЕНТ А ООО', 'Артикул': 'A1', 'Код товара': 'CODE-1',
            'Год': 2026, 'Номер месяца': 9, 'Факт, шт': 5,
        }])
        result = merge_plans_with_1c(plans, actual)
        self.assertEqual(result.iloc[0]['Факт, шт'], 5)

    def test_money_and_negative_metrics_never_become_quantity(self):
        for text in ['Продажи Факт, CNY', 'Факт, %', 'AOP негатив, шт', 'Прогноз, юань']:
            self.assertEqual(parse_metric_type(text), 'ignore')

    def test_duplicate_pair_uses_full_key(self):
        plans = plan_rows().iloc[[1]].copy()
        other = plans.copy()
        other['Производственный индекс'] = 'Other'
        plans = pd.concat([plans, other], ignore_index=True)
        actual = pd.DataFrame([{'Клиент': 'Client', 'Артикул': 'A1', 'Год': 2026, 'Номер месяца': 9,
                              'Факт, шт': 7, 'Факт, CNY': 70, 'Класс товара': 'Class', 'Производственный индекс': 'Other'}])
        result = merge_plans_with_1c(plans, actual)
        self.assertEqual(result['Факт, шт'].tolist(), [0, 7])
        self.assertEqual(result['Факт, CNY'].sum(), 70)
        actual = actual.drop(columns=['Класс товара', 'Производственный индекс'])
        result = merge_plans_with_1c(plans, actual)
        self.assertEqual(result['Факт, шт'].sum(), 0)
        self.assertEqual(result.attrs['matching_report']['unmatched_rows'], 1)

    def test_xml_write_preserves_other_month_and_workbook_parts(self):
        plans = plan_rows()
        plans['Факт, шт'] = [111, 7]
        lookup = build_fact_lookup(plans)
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'plan_redko_source.xlsx'
            output = Path(folder) / 'filled.xlsx'
            source_book(source)
            count = update_manager_workbook(source, output, lookup, 'redko', periods={'2026-09'})
            self.assertEqual(count, 1)
            w = load_workbook(output)
            self.assertEqual(w.active['K4'].value, 8)
            self.assertEqual(w.active['O4'].value, 7)
            self.assertEqual(w.active['Q4'].value, '=SUM(K4,O4)')
            w.close()
            with zipfile.ZipFile(source) as old, zipfile.ZipFile(output) as new:
                for part in old.namelist():
                    if part not in {'xl/worksheets/sheet1.xml', 'xl/workbook.xml'}:
                        self.assertEqual(old.read(part), new.read(part))

    def test_historical_duplicate_does_not_block_current_manager_file(self):
        current = plan_rows().iloc[[1]].copy()
        current['Факт, шт'] = 7
        history = plan_rows().iloc[[0]].copy()
        history['Факт, шт'] = 5
        history_duplicate = history.copy()
        final = pd.concat([current, history, history_duplicate], ignore_index=True)
        lookup = build_fact_lookup(final, periods={'2026-09'})
        self.assertEqual(len(lookup), 1)
        self.assertEqual(next(iter(lookup.values())), 7)

    def test_combined_rows_and_invalidation(self):
        with tempfile.TemporaryDirectory() as folder:
            raw = Path(folder) / 'raw'
            raw.mkdir()
            results = {}
            for manager in ['redko', 'izmaylov']:
                source = raw / f'plan_{manager}_source.xlsx'
                source_book(source, manager)
                results[manager] = {'path': source, 'manager_name': manager}
            output = Path(folder) / 'combined.xlsx'
            result = export_combined_report(raw, results, output)
            self.assertEqual(result['rows'], 2)
            w = load_workbook(output, data_only=True)
            self.assertEqual(w.active.max_row, 4)
            self.assertEqual({w.active.cell(row, 1).value for row in [3, 4]}, {'redko', 'izmaylov'})
            headers = [cell.value for cell in w.active[2]]
            col = headers.index('2026-09 Факт, шт') + 1
            self.assertEqual(w.active.cell(3, col).value, 9)
            w.close()
            self.assertTrue(combined_is_current(output, raw))
            source.unlink()
            self.assertFalse(combined_is_current(output, raw))
