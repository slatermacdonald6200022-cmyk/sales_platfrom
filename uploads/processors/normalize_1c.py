import datetime
import os
import re

import pandas as pd
from openpyxl import load_workbook


def normalize_text(val):
    if pd.isna(val):
        return ""
    return str(val).strip().replace('\n', ' ').replace('\r', ' ')


def normalize_article(val):
    if pd.isna(val):
        return ""
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s.strip()


def parse_number(val):
    """Безопасно преобразует число из 1С с обычными и неразрывными пробелами."""
    if pd.isna(val):
        return 0.0
    try:
        return float(str(val).replace('\xa0', '').replace(' ', '').replace(',', '.'))
    except (TypeError, ValueError):
        return 0.0


def normalize_1c_file(file_path):
    """Извлекает количество; денежные суммы выгрузки не используются."""
    if not file_path or not os.path.exists(file_path):
        return pd.DataFrame()

    df_raw = pd.read_excel(file_path, header=None)

    # 1. Поиск периода в шапке. Не используем месяц по умолчанию:
    # неверно распознанная выгрузка не должна молча изменять факты другого месяца.
    header_lines = [
        " ".join(str(x) for x in df_raw.iloc[i].dropna())
        for i in range(min(15, len(df_raw)))
    ]
    period_lines = [line for line in header_lines if 'период' in line.lower()]
    detected_periods = set()
    for line in period_lines or header_lines[:10]:
        for _, month, year in re.findall(r'(\d{2})\.(\d{2})\.(\d{4})', line):
            month_number = int(month)
            if 1 <= month_number <= 12:
                detected_periods.add((int(year), month_number))

    if not detected_periods:
        raise ValueError(
            'Не удалось определить отчётный период выгрузки. '
            'В шапке файла должен быть указан диапазон дат.'
        )
    if len(detected_periods) != 1:
        raise ValueError('В шапке выгрузки указаны даты из разных месяцев.')
    period_year, period_month = detected_periods.pop()

    # 2. Поиск строки заголовков таблицы
    header_idx = None
    for i in range(min(25, len(df_raw))):
        row_str = " ".join([str(x).lower() for x in df_raw.iloc[i].dropna()])
        if ('клиент' in row_str or 'контрагент' in row_str or 'номенклатура' in row_str) and \
           ('количество' in row_str or 'кол-во' in row_str):
            header_idx = i
            break

    if header_idx is None:
        raise ValueError('Не найдена строка заголовков с количеством и номенклатурой.')

    header_depth = 1
    for offset in (1, 2):
        if header_idx + offset >= len(df_raw):
            break
        next_labels = [normalize_text(value).lower() for value in df_raw.iloc[header_idx + offset].dropna()]
        if any(label in {'заказ клиента', 'номенклатура', 'артикул', 'артикул товара', 'код', 'код товара'}
               for label in next_labels):
            header_depth += 1
        else:
            break
    header_row = pd.Series([
        ' '.join(normalize_text(df_raw.iloc[row, col]) for row in range(header_idx, header_idx + header_depth)).strip()
        for col in range(df_raw.shape[1])
    ])
    col_client_item = 0
    col_article = None
    col_product_code = None
    col_qty = None
    extra_columns = {}

    for c in range(df_raw.shape[1]):
        val = str(header_row.iloc[c]).lower()
        normalized_header = val.strip().replace('\n', ' ')
        if normalized_header == 'класс товара':
            extra_columns['Класс товара'] = c
        if normalized_header in {'производственный индекс', 'при'}:
            extra_columns['Производственный индекс'] = c
        if normalized_header in {'код', 'код товара', 'код из 1с', 'код 1с'}:
            col_product_code = c
        if 'артикул' in val:
            col_article = c
        if 'количество' in val or 'кол-во' in val:
            if col_qty is None:
                col_qty = c

    if col_qty is None:
        raise ValueError('Не найден столбец количества.')

    # В отчёте 1С наименование товара, клиент, реализация и заказ находятся
    # в одной колонке, но на разных уровнях иерархии. Настоящий артикул
    # расположен в отдельной колонке «Артикул товара».
    if col_article is None:
        return pd.DataFrame()

    workbook = load_workbook(file_path, read_only=False, data_only=True)
    worksheet = workbook.active
    has_outline_indents = any(
        (worksheet.cell(row=i + 1, column=col_client_item + 1).alignment.indent or 0) >= 2
        for i in range(header_idx + 1, min(len(df_raw), header_idx + 100))
    )

    rows_data = df_raw.iloc[header_idx + header_depth:].copy()
    records = []
    current_article = None
    current_product_code = None
    current_product = None
    current_extra = {}
    current_client = None

    # В новом отчёте каждая товарная строка уже содержит артикул, код и количество,
    # а клиент находится уровнем выше. Старый формат (товар -> клиент) сохраняется ниже.
    new_item_layout = col_product_code is not None and any(
        normalize_article(row.iloc[col_article]) and normalize_article(row.iloc[col_product_code])
        for _, row in rows_data.head(200).iterrows()
    )

    for row_idx, row in rows_data.iterrows():
        first_cell = normalize_text(row.iloc[col_client_item])
        if not first_cell or 'итого' in first_cell.lower():
            continue

        indent = worksheet.cell(
            row=int(row_idx) + 1,
            column=col_client_item + 1
        ).alignment.indent or 0

        article = normalize_article(row.iloc[col_article])
        product_code = normalize_article(row.iloc[col_product_code]) if col_product_code is not None else ''

        if new_item_layout:
            if indent == 0 and not article and not product_code:
                low = first_cell.lower()
                if not any(marker in low for marker in ('параметр', 'валовая прибыль', 'заказ клиента', 'реализация')):
                    current_client = first_cell
                continue
            if not (article and product_code):
                continue
            qty_val = parse_number(row.iloc[col_qty])
            if not current_client or qty_val == 0:
                continue
            records.append({
                **{name: normalize_article(row.iloc[col]) for name, col in extra_columns.items()},
                'Клиент': current_client,
                'Артикул': article,
                'Код товара': product_code,
                'Наименование': first_cell,
                'Год': period_year,
                'Номер месяца': period_month,
                'Месяц': f"{period_year}-{period_month:02d}",
                'Факт, шт': qty_val,
            })
            continue

        if article:
            current_article = article
            current_product_code = product_code
            current_product = first_cell
            current_extra = {name: normalize_article(row.iloc[col]) for name, col in extra_columns.items()}
            continue

        if not current_article:
            continue

        # Уровень 4 — клиент. Уровни 6 и 8 относятся к реализации и заказу,
        # уровень 0 — итоговая группа товаров. Они не должны становиться фактами.
        if has_outline_indents and not (4 <= indent < 6):
            continue

        # Запасная проверка для экспортов без сохранённых отступов Excel.
        if not has_outline_indents:
            low = first_cell.lower()
            service_markers = (
                'реализация', 'заказ клиента', 'отчет комиссионера',
                'продажи без заказа', 'параметр'
            )
            if any(marker in low for marker in service_markers):
                continue

        qty_val = parse_number(row.iloc[col_qty])
        if qty_val == 0:
            continue

        records.append({
            **{name: normalize_article(row.iloc[col]) or current_extra.get(name, '') for name, col in extra_columns.items()},
            'Клиент': first_cell,
            'Артикул': current_article,
            'Код товара': current_product_code or '',
            'Наименование': current_product,
            'Год': period_year,
            'Номер месяца': period_month,
            'Месяц': f"{period_year}-{period_month:02d}",
            'Факт, шт': qty_val,
        })

    workbook.close()

    result = pd.DataFrame(records)
    result.attrs['report_year'] = period_year
    result.attrs['report_month'] = period_month
    result.attrs['report_period'] = f'{period_year:04d}-{period_month:02d}'
    result.attrs['valuation_method'] = 'quantity_price'
    return result


# Алиас для совместимости
parse_1c_fact_file = normalize_1c_file
process_1c_file = normalize_1c_file
