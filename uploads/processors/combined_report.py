"""Общая широкая таблица: строки менеджеров, месяцы в столбцах.

Не копирует годовые итоги и формулы из разнородных книг. Сохраняет товарные
строки, месячные показатели, исходные описательные поля и комментарии.
"""
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import pandas as pd

from .export_manager_facts import find_sheet_layout, clean_key
from .normalize_manager import normalize_text, normalize_article, parse_month_cell, parse_metric_type, price_columns


def source_manifest(raw_dir):
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(Path(raw_dir).glob('plan_*')) if p.suffix.lower() in {'.xlsx', '.xlsm'}}


def combined_is_current(path, raw_dir):
    path = Path(path)
    try:
        manifest = json.loads(path.with_suffix('.json').read_text(encoding='utf-8'))
        return path.is_file() and bool(manifest) and manifest == source_manifest(raw_dir)
    except (OSError, ValueError):
        return False


def export_combined_report(raw_dir, manager_results, destination):
    columns = []
    records = []
    for manager_id, result in sorted(manager_results.items()):
        source = next(Path(raw_dir).glob(f'plan_{manager_id}_*'))
        original = load_workbook(source, data_only=True, read_only=False)
        filled = load_workbook(result['path'], data_only=False, read_only=False)
        try:
            for sheet in original:
                layout = find_sheet_layout(sheet)
                if not layout:
                    continue
                header = layout['header_row']
                header_frame = pd.DataFrame(list(sheet.iter_rows(max_row=header + 1, values_only=True)))
                prices = {col + 1: (year, half) for col, year, half in price_columns(header_frame, header - 1)}
                mapping = {}
                current_month = None
                occurrences = {}
                for col in range(1, sheet.max_column + 1):
                    top = sheet.cell(header, col).value
                    sub = sheet.cell(header + 1, col).value
                    title = normalize_text(top)
                    detected = parse_month_cell(top)
                    if detected:
                        current_month = detected
                    elif title:
                        current_month = None
                    if col in prices:
                        year, half = prices[col]
                        period = str(year) if year else 'Без периода'
                        if half:
                            period += ' — ' + ('I' if half == 1 else 'II') + ' полугодие'
                        key = ('base', '', 'Цена, CNY без НДС — ' + period)
                    elif current_month:
                        metric = parse_metric_type(sub)
                        if metric in {'aop', 'forecast', 'fact'}:
                            label = {'aop': 'AOP, шт', 'forecast': 'Прогноз, шт', 'fact': 'Факт, шт'}[metric]
                        elif 'негатив' in normalize_text(sub).lower():
                            label = 'AOP негатив, шт'
                        elif 'ком' in normalize_text(sub).lower():
                            label = 'Комментарий'
                        else:
                            continue
                        key = ('month', current_month, label)
                    elif title and not any(t in title.lower() for t in ['итого', 'продажи', 'всего', 'год']):
                        above = normalize_text(sheet.cell(header - 1, col).value) if header > 1 else ''
                        # Одноимённые цены различаются годом/полугодием из верхней строки.
                        if 'цена' in title.lower() and above:
                            title += ' — ' + above
                        canonical = {'Рабочее наименование из 1 С': 'Рабочее наименование из 1С',
                                     'Условия платежа': 'Условия оплаты', 'Коментарий': 'Комментарий'}
                        title = canonical.get(title, title)
                        occurrences[title] = occurrences.get(title, 0) + 1
                        if occurrences[title] > 1:
                            title += f' ({occurrences[title]})'
                        key = ('base', '', title)
                    else:
                        continue
                    mapping[col] = key
                    if key not in columns:
                        columns.append(key)
                client = ''
                for row in range(header + 2, sheet.max_row + 1):
                    raw_client = normalize_text(sheet.cell(row, layout['client_col']).value)
                    if raw_client:
                        client = raw_client
                    article = normalize_article(sheet.cell(row, layout['article_col']).value)
                    if not client or not article or article.lower() in {'итого', 'всего', 'nan', 'none'}:
                        continue
                    record = {key: sheet.cell(row, col).value for col, key in mapping.items()}
                    record[('base', '', 'Менеджер')] = result['manager_name']
                    for col, key in mapping.items():
                        if col == layout['client_col']:
                            record[key] = client
                        if key == ('base', '', 'Менеджер'):
                            record[key] = result['manager_name']
                        if key[0] == 'month' and key[2] == 'Факт, шт':
                            updated = filled[sheet.title].cell(row, col)
                            if updated.data_type != 'f':
                                record[key] = updated.value
                    records.append(record)
        finally:
            original.close()
            filled.close()
    if not records:
        raise ValueError('Не найдены товарные строки для общего файла.')
    order = {'AOP, шт': 0, 'AOP негатив, шт': 1, 'Прогноз, шт': 2, 'Факт, шт': 3, 'Комментарий': 4}
    base = [key for key in columns if key[0] == 'base']
    if ('base', '', 'Менеджер') not in base:
        base.append(('base', '', 'Менеджер'))
    monthly = sorted([key for key in columns if key[0] == 'month'], key=lambda k: (k[1], order[k[2]]))
    columns = base + monthly
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Общий план с фактом'
    sheet.append([k[2] if k[0] == 'base' else k[1] for k in columns])
    sheet.append([k[2] if k[0] == 'base' else f'{k[1]} {k[2]}' for k in columns])
    for record in records:
        sheet.append([record.get(key) for key in columns])
    # Значения из исходных файлов — данные, а не исполняемые формулы.
    for row in sheet.iter_rows(min_row=3):
        for cell in row:
            if cell.data_type == 'f':
                cell.data_type = 's'
    for row in sheet.iter_rows(max_row=2):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
    for idx, key in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(idx)].width = 24 if key[0] == 'base' else 16
    sheet.row_dimensions[1].height = 55
    sheet.row_dimensions[2].height = 32
    sheet.freeze_panes = 'L3'
    sheet.auto_filter.ref = f'A2:{get_column_letter(len(columns))}{sheet.max_row}'
    sheet.sheet_view.showGridLines = False
    destination = Path(destination)
    workbook.save(destination)
    workbook.close()
    destination.with_suffix('.json').write_text(json.dumps(source_manifest(raw_dir), ensure_ascii=False), encoding='utf-8')
    return {'path': destination, 'rows': len(records)}
