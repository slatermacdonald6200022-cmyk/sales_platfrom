"""Проверка Excel-файлов до их сохранения и обработки."""

import datetime
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from django.conf import settings
from openpyxl import load_workbook

from .processors.normalize_manager import (
    BASE_COLUMN_ALIASES,
    parse_metric_type,
    parse_month_cell,
    identify_base_columns,
)


DEFAULT_MAX_FILE_SIZE = 25 * 1024 * 1024
DEFAULT_MAX_UNCOMPRESSED_SIZE = 150 * 1024 * 1024
MAX_ROWS = 200_000
MAX_COLUMNS = 500


class ExcelValidationError(ValueError):
    """Ошибка, которую можно безопасно показать пользователю."""


@dataclass
class ExcelValidationResult:
    rows: int
    cols: int
    sheets_count: int
    sheet_names: str
    periods: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def period_label(self):
        if not self.periods:
            return 'не определён'
        if len(self.periods) <= 3:
            return ', '.join(self.periods)
        return f'{self.periods[0]}–{self.periods[-1]} ({len(self.periods)} мес.)'

    def as_cache_data(self):
        return {
            'rows': self.rows,
            'cols': self.cols,
            'sheets_count': self.sheets_count,
            'sheet_names': self.sheet_names,
            'periods': self.period_label,
        }


def _normalized(value):
    return str(value or '').strip().lower().replace('ё', 'е')


def _format_period(period):
    year, month = period.split('-')
    return f'{int(month):02d}.{int(year):04d}'


def _validate_container(uploaded_file):
    suffix = Path(uploaded_file.name).suffix.lower()
    if suffix != '.xlsx':
        raise ExcelValidationError('Допустим только файл Excel в формате .xlsx.')

    if not uploaded_file.size:
        raise ExcelValidationError('Файл пустой.')

    max_size = getattr(settings, 'MAX_EXCEL_UPLOAD_SIZE', DEFAULT_MAX_FILE_SIZE)
    if uploaded_file.size > max_size:
        max_mb = max_size // (1024 * 1024)
        raise ExcelValidationError(f'Размер файла превышает {max_mb} МБ.')

    uploaded_file.seek(0)
    try:
        with zipfile.ZipFile(uploaded_file) as archive:
            names = set(archive.namelist())
            if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
                raise ExcelValidationError('Файл не является корректной книгой Excel.')
            unpacked_size = sum(item.file_size for item in archive.infolist())
            max_unpacked = getattr(
                settings,
                'MAX_EXCEL_UNCOMPRESSED_SIZE',
                DEFAULT_MAX_UNCOMPRESSED_SIZE,
            )
            if unpacked_size > max_unpacked:
                raise ExcelValidationError('Внутренняя структура файла слишком велика.')
    except (zipfile.BadZipFile, OSError) as exc:
        raise ExcelValidationError('Файл повреждён или не является книгой Excel.') from exc
    finally:
        uploaded_file.seek(0)


def _open_workbook(uploaded_file):
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=False)
    except Exception as exc:
        raise ExcelValidationError('Не удалось открыть книгу Excel. Проверьте файл и повторите загрузку.') from exc
    finally:
        uploaded_file.seek(0)

    if not workbook.worksheets:
        workbook.close()
        raise ExcelValidationError('В книге нет листов.')
    for sheet in workbook.worksheets:
        if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS:
            workbook.close()
            raise ExcelValidationError(
                f'Лист «{sheet.title}» превышает допустимый размер таблицы.'
            )
    return workbook


def _find_manager_layout(sheet):
    scan_rows = list(
        sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 16), values_only=True)
    )
    for row_index, row in enumerate(scan_rows):
        row_text = ' '.join(_normalized(value) for value in row if value is not None)
        has_article = any(alias in row_text for alias in BASE_COLUMN_ALIASES['product_article'])
        has_client = any(alias in row_text for alias in BASE_COLUMN_ALIASES['client'])
        if not (has_article and has_client):
            continue

        main_header = row
        sub_header = scan_rows[row_index + 1] if row_index + 1 < len(scan_rows) else ()
        base_columns = identify_base_columns(main_header)

        periods = set()
        metrics = set()
        current_period = None
        for column_index, value in enumerate(main_header):
            detected_period = parse_month_cell(value)
            if detected_period:
                current_period = detected_period
            sub_value = sub_header[column_index] if column_index < len(sub_header) else None
            metric = parse_metric_type(sub_value) or parse_metric_type(value)
            if current_period and metric in {'aop', 'forecast', 'fact'}:
                periods.add(current_period)
                metrics.add(metric)

        return row_index + 1, base_columns, periods, metrics
    return None


def validate_manager_file(uploaded_file):
    _validate_container(uploaded_file)
    workbook = _open_workbook(uploaded_file)
    valid_sheets = []
    all_periods = set()
    all_metrics = set()
    has_price_column = False
    data_rows = 0
    duplicate_keys = set()

    try:
        for sheet in workbook.worksheets:
            layout = _find_manager_layout(sheet)
            if layout is None:
                continue
            header_row, base_columns, periods, metrics = layout
            if 'product_article' not in base_columns or 'client' not in base_columns:
                continue
            valid_sheets.append(sheet.title)
            all_periods.update(periods)
            all_metrics.update(metrics)
            has_price_column = has_price_column or 'price_cny' in base_columns

            article_column = base_columns['product_article']
            client_column = base_columns['client']
            max_required_column = max(article_column, client_column) + 1
            current_client = ''
            seen_keys = set()
            for row in sheet.iter_rows(
                min_row=header_row + 2,
                min_col=1,
                max_col=max_required_column,
                values_only=True,
            ):
                raw_client = row[client_column]
                if raw_client not in (None, ''):
                    current_client = _normalized(raw_client)
                article = _normalized(row[article_column])
                if not current_client or not article or article in {'итого', 'всего'}:
                    continue
                data_rows += 1
                key = (current_client, article)
                if key in seen_keys:
                    duplicate_keys.add((sheet.title, *key))
                seen_keys.add(key)

        if not valid_sheets:
            raise ExcelValidationError(
                'Не найдена таблица с обязательными столбцами «Клиент» и «Артикул».'
            )
        if not all_periods:
            raise ExcelValidationError('Не удалось определить месяцы в таблице менеджера.')
        if 'aop' not in all_metrics:
            raise ExcelValidationError('Не найдены столбцы плана (AOP) по месяцам.')
        if 'forecast' not in all_metrics:
            raise ExcelValidationError('Не найдены столбцы прогноза по месяцам.')
        if not has_price_column:
            raise ExcelValidationError('Не найден столбец цены в юанях.')
        if data_rows == 0:
            raise ExcelValidationError('В таблице нет товарных строк с артикулами.')

        warnings = []
        skipped_sheets = len(workbook.worksheets) - len(valid_sheets)
        if skipped_sheets:
            warnings.append(f'Пропущено листов без таблицы продаж: {skipped_sheets}.')
        if duplicate_keys:
            warnings.append(
                f'Найдены повторяющиеся пары «клиент + артикул»: {len(duplicate_keys)}. '
                'Проверьте, не приведут ли они к повторному учёту.'
            )

        return ExcelValidationResult(
            rows=max(sheet.max_row for sheet in workbook.worksheets),
            cols=max(sheet.max_column for sheet in workbook.worksheets),
            sheets_count=len(workbook.worksheets),
            sheet_names=', '.join(sheet.title for sheet in workbook.worksheets),
            periods=[_format_period(period) for period in sorted(all_periods)],
            warnings=warnings,
        )
    finally:
        workbook.close()
        uploaded_file.seek(0)


def _find_actual_period(rows):
    lines = [
        ' '.join(str(value) for value in row if value is not None)
        for row in rows[:15]
    ]
    period_lines = [line for line in lines if 'период' in line.lower()]
    dates = []
    # Номера заказов в первых строках могут содержать даты других месяцев.
    # При наличии строки «Период» только она определяет месяц выгрузки.
    for line in period_lines or lines[:10]:
        for day, month, year in re.findall(r'(\d{2})\.(\d{2})\.(\d{4})', line):
            try:
                dates.append(datetime.date(int(year), int(month), int(day)))
            except ValueError:
                continue
    if not dates:
        return None
    periods = {(date.year, date.month) for date in dates}
    if len(periods) != 1:
        raise ExcelValidationError('В заголовке указаны даты из разных месяцев.')
    year, month = periods.pop()
    return f'{year:04d}-{month:02d}'


def validate_actual_file(uploaded_file):
    _validate_container(uploaded_file)
    workbook = _open_workbook(uploaded_file)
    try:
        sheet = workbook.active
        header_rows = list(
            sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25), values_only=True)
        )
        period = _find_actual_period(header_rows)
        if period is None:
            raise ExcelValidationError(
                'Не удалось определить отчётный период. В шапке должен быть указан диапазон дат.'
            )

        header_index = None
        found_fields = set()
        for index, row in enumerate(header_rows):
            # В новых отчётах 1С заголовок занимает три строки:
            # «Клиент», затем «Заказ клиента», затем товарные поля.
            window = header_rows[index:index + 3]
            fields = set()
            for offset, window_row in enumerate(window):
                values = [_normalized(value) for value in window_row]
                if any('номенклатур' in value or 'клиент' in value or 'контрагент' in value for value in values):
                    fields.add('entity')
                if any('артикул' in value for value in values):
                    fields.add('article')
                if any('количество' in value or 'кол-во' in value for value in values):
                    fields.add('quantity')
                if any('выручка' in value or 'сумма' in value for value in values):
                    fields.add('amount')
                if {'entity', 'article', 'quantity'}.issubset(fields):
                    header_index = index + offset + 1
                    found_fields = fields
                    break
            if header_index is not None:
                break

        if header_index is None or len(found_fields) < 3:
            raise ExcelValidationError(
                'Не найдены обязательные столбцы: номенклатура, артикул и количество.'
            )
        if sheet.max_row <= header_index:
            raise ExcelValidationError('В файле нет строк с фактическими данными.')

        return ExcelValidationResult(
            rows=sheet.max_row,
            cols=sheet.max_column,
            sheets_count=len(workbook.worksheets),
            sheet_names=', '.join(item.title for item in workbook.worksheets),
            periods=[_format_period(period)],
        )
    finally:
        workbook.close()
        uploaded_file.seek(0)
