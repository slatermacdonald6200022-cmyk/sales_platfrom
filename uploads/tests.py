import os
import tempfile
import json
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from django.test import SimpleTestCase
from django.test import TestCase, override_settings
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook
from openpyxl.styles import Alignment

from accounts.models import Profile
from accounts.permissions import (
    can_manage_files,
    can_view_company_dashboard,
    can_view_final_dataset,
)
from analytics.models import ProcessingRun
from .views import MANAGERS_LIST, get_user_manager, user_can_access_manager
from .processors.normalize_1c import normalize_1c_file
from .processors.snapshot_engine import merge_plans_with_1c
from .validators import ExcelValidationError, validate_actual_file, validate_manager_file


def workbook_upload(workbook, name='test.xlsx'):
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return SimpleUploadedFile(
        name,
        stream.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def valid_manager_upload(name='manager.xlsx'):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'План'
    sheet.append(['Клиент', 'Артикул', 'Наименование', 'Цена, юань', 'Январь 2026', ''])
    sheet.append(['', '', '', '', 'AOP', 'Прогноз'])
    sheet.append(['Клиент А', 'A-100', 'Товар А', 10, 5, 6])
    return workbook_upload(workbook, name)


def invalid_manager_upload(name='invalid.xlsx'):
    workbook = Workbook()
    workbook.active.append(['Произвольная таблица'])
    return workbook_upload(workbook, name)


def actual_upload_with_period(include_period=True, name='actual.xlsx'):
    workbook = Workbook()
    sheet = workbook.active
    if include_period:
        sheet.append(['Период: 01.05.2026 - 31.05.2026'])
    else:
        sheet.append(['Фактические продажи'])
    sheet.append(['Номенклатура', 'Артикул товара', 'Количество', 'Выручка'])
    sheet.append(['Товар А', 'A-100', 5, 1000])
    return workbook_upload(workbook, name)


class Normalize1CFileTests(SimpleTestCase):
    def test_reads_article_and_client_from_hierarchy_levels(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Период: 01.05.2026 - 31.05.2026"])
        sheet.append([])
        sheet.append([
            "Номенклатура", "", "", "", "Артикул товара", "", "Количество", "Выручка с НДС"
        ])

        rows = [
            (["Готовая продукция/Finish good", "", "", "", "", "", 10, 1000], 0),
            (["Товар А", "", "", "", "4494450600", "", 6, 600], 2),
            (["КЛИЕНТ А ООО", "", "", "", "", "", 6, 600], 4),
            (["Реализация товаров и услуг №1", "", "", "", "", "", 6, 600], 6),
            (["Заказ клиента №1", "", "", "", "", "", 6, 600], 8),
        ]
        for values, indent in rows:
            sheet.append(values)
            sheet.cell(sheet.max_row, 1).alignment = Alignment(indent=indent)

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            path = tmp.name
        try:
            workbook.save(path)
            result = normalize_1c_file(path, source_currency='RUB', cny_rate=10)
        finally:
            workbook.close()
            os.unlink(path)

        self.assertEqual(len(result), 1)
        row = result.iloc[0]
        self.assertEqual(row["Клиент"], "КЛИЕНТ А ООО")
        self.assertEqual(row["Артикул"], "4494450600")
        self.assertEqual(row["Наименование"], "Товар А")
        self.assertEqual(row["Год"], 2026)
        self.assertEqual(row["Номер месяца"], 5)
        self.assertEqual(row["Факт, шт"], 6)
        self.assertEqual(row["Факт, CNY"], 60)

    def test_missing_period_is_not_replaced_with_hardcoded_month(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Фактические продажи'])
        sheet.append(['Номенклатура', 'Артикул товара', 'Количество', 'Выручка'])
        sheet.append(['Товар А', 'A-100', 5, 1000])

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            path = tmp.name
        try:
            workbook.save(path)
            with self.assertRaisesMessage(ValueError, 'отчётный период'):
                normalize_1c_file(path, source_currency='CNY')
        finally:
            workbook.close()
            os.unlink(path)


class ExcelUploadValidationTests(SimpleTestCase):
    def test_valid_manager_file_reports_detected_period(self):
        result = validate_manager_file(valid_manager_upload())
        self.assertEqual(result.period_label, '01.2026')
        self.assertEqual(result.sheets_count, 1)

    def test_non_xlsx_file_is_rejected(self):
        uploaded = SimpleUploadedFile('plan.csv', b'client,article')
        with self.assertRaisesMessage(ExcelValidationError, 'формате .xlsx'):
            validate_manager_file(uploaded)

    def test_actual_file_requires_report_period(self):
        with self.assertRaisesMessage(ExcelValidationError, 'отчётный период'):
            validate_actual_file(actual_upload_with_period(include_period=False))

    def test_valid_actual_file_reports_month(self):
        result = validate_actual_file(actual_upload_with_period())
        self.assertEqual(result.period_label, '05.2026')


class PeriodScopedMergeTests(SimpleTestCase):
    def test_current_export_changes_only_its_own_month(self):
        plans = pd.DataFrame([
            {
                'Клиент': 'Клиент А', 'Артикул': 'A-100', 'Наименование': 'Товар А',
                'Менеджер': 'Редько Вадим', 'Поставщик': 'Поставщик',
                'Год': 2026, 'Номер месяца': 5, 'Месяц': 'Май',
                'AOP, шт': 10, 'Прогноз, шт': 9,
                'Цена, юань, без НДС 1 п/г 2026': 10,
            },
            {
                'Клиент': 'Клиент А', 'Артикул': 'A-100', 'Наименование': 'Товар А',
                'Менеджер': 'Редько Вадим', 'Поставщик': 'Поставщик',
                'Год': 2026, 'Номер месяца': 6, 'Месяц': 'Июнь',
                'AOP, шт': 20, 'Прогноз, шт': 18,
                'Цена, юань, без НДС 1 п/г 2026': 10,
            },
        ])
        actuals = pd.DataFrame([{
            'Клиент': 'Клиент А', 'Артикул': 'A-100',
            'Год': 2026, 'Номер месяца': 5,
            'Факт, шт': 7, 'Факт, CNY': 70,
        }])
        history = pd.DataFrame([
            {
                '_key_client': 'клиент а', '_key_article': 'a-100',
                '_key_month': '2026-05', 'Факт, шт': 4, 'Факт, CNY': 40,
            },
            {
                '_key_client': 'клиент а', '_key_article': 'a-100',
                '_key_month': '2026-06', 'Факт, шт': 6, 'Факт, CNY': 60,
            },
        ])

        result = merge_plans_with_1c(plans, actuals, history)
        may = result[result['Номер месяца'] == 5].iloc[0]
        june = result[result['Номер месяца'] == 6].iloc[0]

        self.assertEqual((may['Факт, шт'], may['Факт, CNY']), (7, 70))
        self.assertEqual((june['Факт, шт'], june['Факт, CNY']), (6, 60))
        self.assertEqual(len(result.columns), 15)
        report = result.attrs['matching_report']
        self.assertEqual(report['matched_rows'], 1)
        self.assertEqual(report['unmatched_rows'], 0)

    def test_unmatched_fact_is_reported_without_entering_result(self):
        plans = pd.DataFrame([{
            'Клиент': 'Клиент А', 'Артикул': 'A-100', 'Наименование': 'Товар А',
            'Менеджер': 'Редько Вадим', 'Поставщик': '',
            'Год': 2026, 'Номер месяца': 5, 'Месяц': 'Май',
            'AOP, шт': 10, 'Прогноз, шт': 9,
            'Цена, юань, без НДС 1 п/г 2026': 10,
        }])
        actuals = pd.DataFrame([{
            'Клиент': 'Неизвестный клиент', 'Артикул': 'X-1',
            'Год': 2026, 'Номер месяца': 5,
            'Факт, шт': 7, 'Факт, CNY': 70,
        }])

        result = merge_plans_with_1c(plans, actuals)
        report = result.attrs['matching_report']

        self.assertEqual(report['matched_rows'], 0)
        self.assertEqual(report['unmatched_rows'], 1)
        self.assertEqual(report['unmatched_amount_cny'], 70)
        self.assertEqual(
            report['unmatched_df'].iloc[0]['Причина'],
            'Клиент не найден в планах',
        )
        self.assertEqual(result.iloc[0]['Факт, CNY'], 0)


class RoleAccessTests(TestCase):
    def create_user(self, username, role, manager_name=''):
        user = User.objects.create_user(username=username, password='StrongPass2026!')
        Profile.objects.update_or_create(
            user=user,
            defaults={'role': role, 'manager_name': manager_name},
        )
        user.refresh_from_db()
        return user

    def setUp(self):
        self.manager = self.create_user('redko', 'manager', 'Редько Вадим')
        self.other_manager = self.create_user('khusnutdinov', 'manager', 'Хуснутдинов')
        self.analyst = self.create_user('analyst', 'analyst')
        self.director = self.create_user('director', 'director')

    def test_role_matrix(self):
        self.assertFalse(can_manage_files(self.manager))
        self.assertFalse(can_view_company_dashboard(self.manager))
        self.assertFalse(can_view_final_dataset(self.manager))

        self.assertTrue(can_manage_files(self.analyst))
        self.assertTrue(can_view_final_dataset(self.analyst))

        self.assertFalse(can_manage_files(self.director))
        self.assertTrue(can_view_company_dashboard(self.director))
        self.assertFalse(can_view_final_dataset(self.director))

    def test_manager_is_resolved_by_exact_account_assignment(self):
        assigned = get_user_manager(self.manager)
        self.assertIsNotNone(assigned)
        self.assertEqual(assigned['id'], 'redko')

        redko = next(item for item in MANAGERS_LIST if item['id'] == 'redko')
        khusnutdinov = next(item for item in MANAGERS_LIST if item['id'] == 'khusnutdinov')
        self.assertTrue(user_can_access_manager(self.manager, redko))
        self.assertFalse(user_can_access_manager(self.manager, khusnutdinov))

    def test_manager_cannot_download_other_manager_report(self):
        self.client.force_login(self.manager)
        response = self.client.get(
            reverse('download_manager_report', kwargs={'manager_id': 'khusnutdinov'})
        )
        self.assertEqual(response.status_code, 403)

    def test_manager_cannot_download_final_dataset(self):
        self.client.force_login(self.manager)
        response = self.client.get(reverse('download_final_excel'))
        self.assertEqual(response.status_code, 403)

    def test_director_cannot_open_upload_page(self):
        self.client.force_login(self.director)
        response = self.client.get(reverse('upload_files'))
        self.assertEqual(response.status_code, 403)

    def test_manager_cannot_start_processing(self):
        self.client.force_login(self.manager)
        response = self.client.post(reverse('run_etl_api'))
        self.assertEqual(response.status_code, 403)

    def test_file_deletion_rejects_get_request(self):
        self.client.force_login(self.analyst)
        response = self.client.get(
            reverse('delete_file', kwargs={'file_type': 'manager', 'target_id': 'redko'})
        )
        self.assertEqual(response.status_code, 405)

    def test_processing_page_hides_final_dataset_from_manager(self):
        self.client.force_login(self.other_manager)
        response = self.client.get(reverse('processing_page'))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Скачать итоговый файл')

    def test_processing_history_is_limited_for_manager(self):
        visible = ProcessingRun.objects.create(
            user=self.analyst,
            status=ProcessingRun.STATUS_SUCCESS,
            report_year=2026,
            report_month=5,
            manager_reports={'redko': 'data/processed/redko.xlsx'},
        )
        hidden = ProcessingRun.objects.create(
            user=self.analyst,
            status=ProcessingRun.STATUS_SUCCESS,
            report_year=2026,
            report_month=6,
            manager_reports={'khusnutdinov': 'data/processed/kh.xlsx'},
        )

        self.client.force_login(self.manager)
        response = self.client.get(reverse('processing_history'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse('processing_run_detail', args=[visible.pk]))
        self.assertNotContains(response, reverse('processing_run_detail', args=[hidden.pk]))

    def test_director_cannot_download_processing_files(self):
        run = ProcessingRun.objects.create(
            user=self.analyst,
            status=ProcessingRun.STATUS_SUCCESS,
            final_file='data/processed/final.xlsx',
        )
        self.client.force_login(self.director)
        response = self.client.get(
            reverse('download_processing_file', args=[run.pk, 'final'])
        )
        self.assertEqual(response.status_code, 403)

    def test_manager_dashboard_contains_only_assigned_manager_data(self):
        rows = [
            {
                'AOP, CNY': 100.0,
                'Прогноз, CNY': 90.0,
                'Факт, CNY': 80.0,
                'AOP, шт': 10.0,
                'Прогноз, шт': 9.0,
                'Факт, шт': 8.0,
                'Год': 2026,
                'Номер месяца': 5,
                'Месяц': 'Май',
                'Менеджер': 'Редько Вадим',
                'Клиент': 'Клиент Редько',
                'Артикул': 'R-1',
                'Наименование': 'Товар Редько',
            },
            {
                'AOP, CNY': 999999.0,
                'Прогноз, CNY': 999999.0,
                'Факт, CNY': 999999.0,
                'AOP, шт': 1.0,
                'Прогноз, шт': 1.0,
                'Факт, шт': 1.0,
                'Год': 2026,
                'Номер месяца': 5,
                'Месяц': 'Май',
                'Менеджер': 'Хуснутдинов А.',
                'Клиент': 'Другой клиент',
                'Артикул': 'K-1',
                'Наименование': 'Другой товар',
            },
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            final_path = Path(temp_dir) / 'FINAL_SALES_FACT_TABLE.xlsx'
            pd.DataFrame(rows).to_excel(final_path, index=False)

            self.client.force_login(self.manager)
            with patch('accounts.views.get_latest_final_file', return_value=final_path):
                response = self.client.get(reverse('dashboard'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(json.loads(response.context['dyn_managers']), ['Редько Вадим'])
        self.assertEqual(response.context['kpi']['month_fact'], '80 ¥')

    def test_dashboard_uses_period_saved_by_latest_processing(self):
        rows = [
            {
                'AOP, CNY': 100.0, 'Прогноз, CNY': 90.0, 'Факт, CNY': 80.0,
                'AOP, шт': 10.0, 'Прогноз, шт': 9.0, 'Факт, шт': 8.0,
                'Год': 2026, 'Номер месяца': 5, 'Месяц': 'Май',
                'Менеджер': 'Редько Вадим', 'Клиент': 'Клиент А',
                'Артикул': 'A-1', 'Наименование': 'Товар А',
            },
            {
                'AOP, CNY': 200.0, 'Прогноз, CNY': 190.0, 'Факт, CNY': 170.0,
                'AOP, шт': 20.0, 'Прогноз, шт': 19.0, 'Факт, шт': 17.0,
                'Год': 2027, 'Номер месяца': 1, 'Месяц': 'Январь',
                'Менеджер': 'Редько Вадим', 'Клиент': 'Клиент А',
                'Артикул': 'A-1', 'Наименование': 'Товар А',
            },
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            final_path = Path(temp_dir) / 'FINAL_SALES_FACT_TABLE.xlsx'
            pd.DataFrame(rows).to_excel(final_path, index=False)
            final_path.with_name('processing_metadata.json').write_text(
                json.dumps({'report_year': 2027, 'report_month': 1}),
                encoding='utf-8',
            )

            self.client.force_login(self.manager)
            with patch('accounts.views.get_latest_final_file', return_value=final_path):
                response = self.client.get(reverse('dashboard'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['selected_period'], 'reporting_period')
        self.assertEqual(response.context['target_year'], 2027)
        self.assertEqual(response.context['kpi']['month_fact'], '170 ¥')

    def test_invalid_upload_keeps_previous_manager_file(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            raw_dir = Path(temp_dir) / 'data' / 'raw'
            raw_dir.mkdir(parents=True)
            old_file = raw_dir / 'plan_redko_previous.xlsx'
            old_file.write_bytes(b'previous-file')

            self.client.force_login(self.analyst)
            with override_settings(BASE_DIR=Path(temp_dir)):
                response = self.client.post(
                    reverse('upload_files'),
                    {'file_manager_redko': invalid_manager_upload()},
                )

            self.assertEqual(response.status_code, 302)
            self.assertTrue(old_file.exists())
            self.assertEqual(old_file.read_bytes(), b'previous-file')

    def test_valid_upload_replaces_previous_manager_file(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            raw_dir = Path(temp_dir) / 'data' / 'raw'
            raw_dir.mkdir(parents=True)
            old_file = raw_dir / 'plan_redko_previous.xlsx'
            old_file.write_bytes(b'previous-file')

            self.client.force_login(self.analyst)
            with override_settings(BASE_DIR=Path(temp_dir)):
                response = self.client.post(
                    reverse('upload_files'),
                    {'file_manager_redko': valid_manager_upload('new_plan.xlsx')},
                )

            saved_files = list(raw_dir.glob('plan_redko_*.xlsx'))
            self.assertEqual(response.status_code, 302)
            self.assertEqual(len(saved_files), 1)
            self.assertEqual(saved_files[0].name, 'plan_redko_new_plan.xlsx')
            self.assertFalse(old_file.exists())

# Create your tests here.
