import os
import datetime
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.conf import settings

# Реестр 10 закрепленных менеджеров
MANAGERS_LIST = [
    {'id': 'tsarev', 'name': 'Царев Михаил', 'username': 'tsarev'},
    {'id': 'khusnutdinov', 'name': 'Хуснутдинов', 'username': 'khusnutdinov'},
    {'id': 'redko', 'name': 'Редько', 'username': 'redko'},
    {'id': 'izmaylov', 'name': 'Измайлов', 'username': 'izmaylov'},
    {'id': 'mustafin', 'name': 'Мустафин', 'username': 'mustafin'},
    {'id': 'polyakov', 'name': 'Поляков', 'username': 'polyakov'},
    {'id': 'prasolov', 'name': 'Прасолов Н. / Соловьёв В.', 'username': 'prasolov'},
    {'id': 'fomichev', 'name': 'Фомичев Владимир', 'username': 'fomichev'},
    {'id': 'khoroshevsky', 'name': 'Хорошевский Александр', 'username': 'khoroshevsky'},
    {'id': 'ushakov', 'name': 'Ушаков Алексей', 'username': 'ushakov'},
]


def check_is_admin(user):
    """Проверка прав администратора/аналитика"""
    if user.is_superuser or user.is_staff:
        return True
    profile = getattr(user, 'profile', None)
    return bool(profile and profile.role in ['analyst', 'director'])


def get_excel_metadata(file_path):
    """Считывает расширенные метаданные Excel-файла (размер, строки, колонки, листы)"""
    if not os.path.exists(file_path):
        return None

    stat = os.stat(file_path)
    modified_time = datetime.datetime.fromtimestamp(stat.st_mtime).strftime('%d.%m.%Y %H:%M')
    size_kb = round(stat.st_size / 1024, 1)
    size_str = f"{round(size_kb / 1024, 2)} МБ" if size_kb > 1024 else f"{size_kb} КБ"

    rows = 'Н/Д'
    cols = 'Н/Д'
    sheets_count = 1
    sheet_names = ''

    # Быстрое чтение структуры книги без загрузки всего массива в память
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheets = wb.sheetnames
        sheets_count = len(sheets)
        sheet_names = ", ".join(sheets)
        active_sheet = wb[sheets[0]]
        rows = active_sheet.max_row or 'Н/Д'
        cols = active_sheet.max_column or 'Н/Д'
        wb.close()
    except Exception:
        try:
            import pandas as pd
            xl = pd.ExcelFile(file_path)
            sheets_count = len(xl.sheet_names)
            sheet_names = ", ".join(xl.sheet_names)
            df_preview = pd.read_excel(file_path, nrows=5)
            cols = len(df_preview.columns)
        except Exception:
            pass

    return {
        'size': size_str,
        'modified': modified_time,
        'rows': rows,
        'cols': cols,
        'sheets_count': sheets_count,
        'sheet_names': sheet_names,
    }


def scan_raw_directory(upload_dir):
    """Сканирует папку data/raw/ и формирует реестр прикрепленных файлов"""
    if not os.path.exists(upload_dir):
        return {}, None

    files = os.listdir(upload_dir)
    manager_files = {}
    fact_1c_info = None

    for fname in files:
        fpath = os.path.join(upload_dir, fname)
        if not os.path.isfile(fpath):
            continue

        fname_lower = fname.lower()
        if not (fname_lower.endswith('.xlsx') or fname_lower.endswith('.xls')):
            continue

        # 1. Поиск выгрузки 1C ERP
        if fname.startswith('fact_1c_') or '1c' in fname_lower or 'факт' in fname_lower:
            orig_name = fname.replace('fact_1c_', '', 1) if fname.startswith('fact_1c_') else fname
            meta = get_excel_metadata(fpath) or {}
            fact_1c_info = {
                'stored_filename': fname,
                'display_name': orig_name,
                **meta
            }
            continue

        # 2. Поиск планов менеджеров
        for mgr in MANAGERS_LIST:
            mgr_id = mgr['id']
            prefix = f"plan_{mgr_id}_"
            surname = mgr['name'].lower().split()[0]
            if fname.startswith(prefix) or surname in fname_lower:
                orig_name = fname.replace(prefix, '', 1) if fname.startswith(prefix) else fname
                meta = get_excel_metadata(fpath) or {}
                manager_files[mgr_id] = {
                    'stored_filename': fname,
                    'display_name': orig_name,
                    **meta
                }
                break

    return manager_files, fact_1c_info


def remove_manager_file_from_disk(upload_dir, mgr_id):
    """Удаляет файл конкретного менеджера с диска"""
    prefix = f"plan_{mgr_id}_"
    if not os.path.exists(upload_dir):
        return False
    deleted = False
    for fname in os.listdir(upload_dir):
        if fname.startswith(prefix):
            try:
                os.remove(os.path.join(upload_dir, fname))
                deleted = True
            except OSError:
                pass
    return deleted


@login_required
def upload_view(request):
    """
    Модуль 1: Загрузка файлов с автосохранением и удалением
    """
    user = request.user
    is_admin = check_is_admin(user)
    upload_dir = os.path.join(settings.BASE_DIR, 'data', 'raw')
    os.makedirs(upload_dir, exist_ok=True)

    # Обработка автоматической загрузки через POST
    if request.method == 'POST':
        saved_count = 0

        # Загрузка файла 1С:ERP
        if is_admin and 'file_1c' in request.FILES:
            f_1c = request.FILES['file_1c']
            for fname in os.listdir(upload_dir):
                if fname.startswith('fact_1c_'):
                    try:
                        os.remove(os.path.join(upload_dir, fname))
                    except OSError:
                        pass

            clean_name = f_1c.name.replace(' ', '_')
            dest_name = f"fact_1c_{clean_name}"
            dest_path = os.path.join(upload_dir, dest_name)
            with open(dest_path, 'wb+') as dest:
                for chunk in f_1c.chunks():
                    dest.write(chunk)
            messages.success(request, f'Выгрузка 1С «{f_1c.name}» успешно сохранена.')
            return redirect('upload_files')

        # Загрузка планов менеджеров
        for mgr in MANAGERS_LIST:
            field_name = f"file_manager_{mgr['id']}"
            if field_name in request.FILES:
                # Проверка прав: менеджер может загрузить только свой слот
                if not is_admin:
                    user_matched = (
                        mgr['username'] == user.username or
                        mgr['id'] in user.username.lower() or
                        (hasattr(user, 'profile') and user.profile.manager_name and mgr['name'].startswith(user.profile.manager_name.split()[0]))
                    )
                    if not user_matched:
                        continue

                m_file = request.FILES[field_name]
                remove_manager_file_from_disk(upload_dir, mgr['id'])

                clean_name = m_file.name.replace(' ', '_')
                dest_name = f"plan_{mgr['id']}_{clean_name}"
                dest_path = os.path.join(upload_dir, dest_name)

                with open(dest_path, 'wb+') as dest:
                    for chunk in m_file.chunks():
                        dest.write(chunk)
                messages.success(request, f'Файл «{m_file.name}» для {mgr["name"]} успешно сохранен.')
                return redirect('upload_files')

    # Получаем актуальное состояние файлов
    manager_files_map, fact_1c_info = scan_raw_directory(upload_dir)

    if is_admin:
        target_managers = MANAGERS_LIST
    else:
        target_managers = [
            m for m in MANAGERS_LIST
            if m['username'] == user.username or m['id'] in user.username.lower() or (
                hasattr(user, 'profile') and user.profile.manager_name and m['name'].startswith(user.profile.manager_name.split()[0])
            )
        ]
        if not target_managers:
            mgr_title = getattr(getattr(user, 'profile', None), 'manager_name', None) or user.get_full_name() or user.username
            target_managers = [{'id': user.username, 'name': mgr_title, 'username': user.username}]

    slots = []
    for mgr in target_managers:
        f_info = manager_files_map.get(mgr['id'])
        slots.append({
            'id': mgr['id'],
            'name': mgr['name'],
            'username': mgr['username'],
            'is_uploaded': bool(f_info),
            'info': f_info,
        })

    context = {
        'is_admin': is_admin,
        'slots': slots,
        'total_managers_count': len(MANAGERS_LIST),
        'uploaded_managers_count': len(manager_files_map),
        'fact_1c': fact_1c_info,
    }
    return render(request, 'uploads/upload.html', context)


@login_required
def delete_file_view(request, file_type, target_id):
    """Удаление файла менеджера или файла 1С"""
    user = request.user
    is_admin = check_is_admin(user)
    upload_dir = os.path.join(settings.BASE_DIR, 'data', 'raw')

    if file_type == 'fact_1c':
        if not is_admin:
            messages.error(request, 'У вас нет прав для удаления выгрузки 1С.')
            return redirect('upload_files')
        for fname in os.listdir(upload_dir):
            if fname.startswith('fact_1c_') or fname == target_id:
                try:
                    os.remove(os.path.join(upload_dir, fname))
                except OSError:
                    pass
        messages.success(request, 'Файл выгрузки 1С успешно удален.')

    elif file_type == 'manager':
        # Проверка прав: менеджер может удалить только свой файл
        if not is_admin:
            can_delete = (
                target_id == user.username or
                target_id in user.username.lower()
            )
            if not can_delete:
                messages.error(request, 'Вы можете удалить только свой собственный файл.')
                return redirect('upload_files')

        remove_manager_file_from_disk(upload_dir, target_id)
        messages.success(request, 'Файл плана успешно удален из хранилища.')

    return redirect('upload_files')


@login_required
def readiness_view(request):
    """Модуль 2: Контроль готовности и запуск ETL"""
    return render(request, 'uploads/processing_readiness.html')


@login_required
def compare_view(request):
    """Модуль сравнения срезов"""
    return render(request, 'uploads/compare.html')